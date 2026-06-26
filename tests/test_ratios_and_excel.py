"""Ratio engine + live-formula Excel guarantees (acceptance criteria 3 & 4)."""

from pathlib import Path

import pytest
from openpyxl import load_workbook

from ivanalysis import build_workbook, company_from_dict, generate_report
from ivanalysis.metrics import run_ratios


def _by_name(results):
    return {r.name: r for r in results}


@pytest.fixture
def company():
    return company_from_dict({
        "name": "TestCo",
        "periods": [
            {"label": "FY23", "fiscal_year": 2023,
             "revenue": 1000, "cogs": 600, "gross_profit": 400,
             "ebit": 200, "interest_expense": 40, "net_income": 120,
             "current_assets": 500, "inventory": 150, "current_liabilities": 250,
             "total_assets": 1200, "net_fixed_assets": 700,
             "total_liabilities": 600, "total_equity": 600,
             "long_term_debt": 300, "accounts_receivable": 100,
             "diluted_shares": 100},
            {"label": "FY24", "fiscal_year": 2024,
             "revenue": 1200, "cogs": 700, "gross_profit": 500,
             "ebit": 260, "interest_expense": 40, "net_income": 160,
             "current_assets": 600, "inventory": 180, "current_liabilities": 250,
             "total_assets": 1400, "net_fixed_assets": 800,
             "total_liabilities": 650, "total_equity": 750,
             "long_term_debt": 320, "accounts_receivable": 120,
             "diluted_shares": 100},
        ],
        "assumptions": {"price_per_share": 24, "book_value_per_share": 7.5,
                        "tax_rate": 0.21, "risk_free_rate": 0.04,
                        "market_return": 0.09, "beta": 1.1,
                        "dividend_per_share": 0.5, "dividend_growth_g": 0.03,
                        "required_return_k": 0.09, "cost_of_debt_pretax": 0.06},
    })


def test_ratio_values(company):
    m = _by_name(run_ratios(company, 1))  # FY24
    assert abs(m["Current Ratio"].value - 600 / 250) < 1e-9
    assert abs(m["Quick (Acid-Test) Ratio"].value - (600 - 180) / 250) < 1e-9
    assert abs(m["Net Working Capital"].value - 350) < 1e-9
    assert abs(m["Debt Ratio"].value - 650 / 1400) < 1e-9
    assert abs(m["Times Interest Earned"].value - 260 / 40) < 1e-9
    assert abs(m["Gross Profit Margin"].value - 500 / 1200) < 1e-9
    assert abs(m["Net Profit Margin"].value - 160 / 1200) < 1e-9
    assert abs(m["Return on Equity (ROE)"].value - 160 / 750) < 1e-9
    assert abs(m["Equity Multiplier"].value - 1400 / 750) < 1e-9


def test_dupont_identity(company):
    m = _by_name(run_ratios(company, 1))
    roe = m["Return on Equity (ROE)"].value
    dupont = m["ROE (DuPont 3-step)"].value
    assert abs(roe - dupont) < 1e-9  # decomposition reconciles to direct ROE


def test_inventory_turnover_uses_average(company):
    m = _by_name(run_ratios(company, 1))
    avg_inv = (180 + 150) / 2
    assert abs(m["Inventory Turnover"].value - 700 / avg_inv) < 1e-9
    assert "AVERAGE(" in m["Inventory Turnover"].formula


def test_missing_input_is_na_not_error():
    c = company_from_dict({"name": "Sparse", "periods": [
        {"label": "FY24", "fiscal_year": 2024, "revenue": 100, "net_income": 10}]})
    m = _by_name(run_ratios(c, 0))
    assert m["Current Ratio"].value is None        # no current assets/liabs
    assert m["Current Ratio"].display() == "n/a"
    assert m["Net Profit Margin"].value == 0.1     # this one still computes


def test_formula_references_named_ranges(company):
    m = _by_name(run_ratios(company, 1))
    f = m["Current Ratio"].formula
    assert f == "=CurrentAssets_FY24/CurrentLiab_FY24"


def test_excel_has_live_formulas(company, tmp_path):
    out = tmp_path / "test.xlsx"
    build_workbook(company, out)
    assert out.exists()
    wb = load_workbook(out)  # data_only=False -> formulas preserved
    assert {"Inputs", "Assumptions", "Ratios", "Valuation", "CapBudget", "Charts"} <= set(wb.sheetnames)

    # A ratio cell must be a live formula, not a baked number.
    ratios = wb["Ratios"]
    formula_cells = [c for row in ratios.iter_rows() for c in row
                     if isinstance(c.value, str) and c.value.startswith("=")]
    assert formula_cells, "Ratios sheet has no live formulas"
    assert any("CurrentAssets_FY24" in c.value or "/" in c.value for c in formula_cells)

    # Named ranges must exist so the formulas resolve.
    assert "CurrentAssets_FY24" in wb.defined_names
    assert "TaxRate" in wb.defined_names

    # CapBudget must carry NPV/IRR formulas.
    cap = wb["CapBudget"]
    cap_formulas = " ".join(str(c.value) for row in cap.iter_rows() for c in row
                            if isinstance(c.value, str))
    assert "NPV(" in cap_formulas and "IRR(" in cap_formulas


def test_report_generates_with_verdict(company):
    report = generate_report(company)
    assert "Investment-Viability Report" in report
    assert "Verdict:" in report
    assert any(v in report for v in ("VIABLE", "CONDITIONAL", "NOT VIABLE"))
    assert "not investment, tax, or legal advice" in report.lower()
